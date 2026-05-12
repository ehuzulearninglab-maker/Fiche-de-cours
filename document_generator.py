"""
Document generator - Creates Word (.docx) and PDF files matching the official canevas.
"""
import re
import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import subprocess


def set_cell_border(cell, **kwargs):
    """Set cell border properties."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            element = OxmlElement(f'w:{edge}')
            for key in ['sz', 'val', 'color', 'space']:
                if key in edge_data:
                    element.set(qn(f'w:{key}'), str(edge_data[key]))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)


def parse_fiche_content(content: str) -> dict:
    """Parse the LLM-generated fiche content into structured sections."""
    sections = {
        "matiere": "",
        "sa": "",
        "seance": "",
        "cours": "",
        "date": "",
        "titre": "",
        "duree": "",
        "competences_disc": "",
        "comp_transv": "",
        "connaissances_tech": "",
        "strategie": "",
        "materiel": "",
        "deroulement": [],
        "resume": "",
        "notes": "",
    }

    lines = content.strip().split("\n")
    current_section = None
    deroulement_rows = []

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # Parse header fields
        if line.upper().startswith("MATIÈRE") or line.upper().startswith("MATIERE"):
            sections["matiere"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif "SA N°" in line.upper() or "SAN°" in line.upper():
            sections["sa"] = line.split(":", 1)[-1].strip() if ":" in line else line
        elif line.upper().startswith("SÉANCE") or line.upper().startswith("SEANCE"):
            sections["seance"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif line.upper().startswith("COURS"):
            sections["cours"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif line.upper().startswith("TITRE"):
            sections["titre"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif line.upper().startswith("DURÉE") or line.upper().startswith("DUREE"):
            sections["duree"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif "COMPÉTENCES DISCIPLINAIRES" in line.upper() or "COMPETENCES DISCIPLINAIRES" in line.upper():
            sections["competences_disc"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif "COMP. TRANSV" in line.upper():
            sections["comp_transv"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif "CONNAISSANCES TECHNIQUES" in line.upper():
            sections["connaissances_tech"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif "STRATÉGIE" in line.upper() or "STRATEGIE" in line.upper():
            sections["strategie"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif "MATÉRIEL" in line.upper() or "MATERIEL" in line.upper():
            sections["materiel"] = line.split(":", 1)[-1].strip() if ":" in line else ""
        elif line.upper().startswith("RÉSUMÉ") or line.upper().startswith("RESUME"):
            resume_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().upper().startswith("NOTES"):
                resume_lines.append(lines[i])
                i += 1
            sections["resume"] = "\n".join(resume_lines).strip()
            continue
        elif line.upper().startswith("NOTES PERSONNELLES"):
            notes_lines = []
            i += 1
            while i < len(lines):
                notes_lines.append(lines[i])
                i += 1
            sections["notes"] = "\n".join(notes_lines).strip()
            continue
        elif "|" in line and ("CONSIGNE" in line.upper() or "RÉSULTAT" in line.upper() or "RESULTAT" in line.upper()):
            # Table header - skip
            i += 1
            # Skip separator line
            if i < len(lines) and set(lines[i].strip().replace(" ", "")) <= set("|:-"):
                i += 1
            continue
        elif "|" in line:
            # Table row
            raw_parts = line.split("|")
            # Drop only the leading/trailing empty strings produced by surrounding pipes,
            # but keep empty middle cells (otherwise rows like `| **SECTION** | |`
            # collapse to a single non-empty part and the row is silently dropped).
            if raw_parts and raw_parts[0].strip() == "":
                raw_parts = raw_parts[1:]
            if raw_parts and raw_parts[-1].strip() == "":
                raw_parts = raw_parts[:-1]
            parts = [p.strip() for p in raw_parts]
            if len(parts) >= 1:
                consigne = parts[0]
                resultat = parts[1] if len(parts) > 1 else ""
                # Markdown-bold-only consigne with empty resultat = section header row
                is_section_header = (
                    not resultat
                    and consigne.startswith("**")
                    and consigne.endswith("**")
                )
                row = {
                    "consigne": consigne.strip("*").strip() if is_section_header else consigne,
                    "resultat": resultat,
                }
                if is_section_header:
                    row["is_header"] = True
                deroulement_rows.append(row)
        elif line.startswith("**") and line.endswith("**"):
            # Section header in deroulement
            section_name = line.strip("*").strip()
            deroulement_rows.append({
                "consigne": section_name,
                "resultat": "",
                "is_header": True
            })

        i += 1

    sections["deroulement"] = deroulement_rows
    return sections


def generate_docx(content: str, output_path: str, raw_mode: bool = False) -> str:
    """Generate a Word document from fiche content."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)

    if raw_mode:
        # Simple mode: just dump the content
        _generate_raw_docx(doc, content)
    else:
        # Try structured parsing
        parsed = parse_fiche_content(content)
        if parsed["deroulement"]:
            _generate_structured_docx(doc, parsed)
        else:
            _generate_raw_docx(doc, content)

    doc.save(output_path)
    return output_path


def _generate_structured_docx(doc, parsed):
    """Generate a properly structured fiche document."""
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"FICHE DE PREPARATION : {parsed['matiere']}")
    run.bold = True
    run.font.size = Pt(13)
    run.font.name = 'Times New Roman'

    # Header info
    header_fields = [
        f"SA N° : {parsed['sa']}    SÉANCE : {parsed['seance']}    Cours : {parsed['cours']}    Date : {parsed['date'] or '……………'}",
        f"TITRE : {parsed['titre']}",
        f"Durée : {parsed['duree'] or '60 min'}",
    ]
    for field in header_fields:
        p = doc.add_paragraph()
        run = p.add_run(field)
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)

    # Planning section
    planning_title = doc.add_paragraph()
    run = planning_title.add_run("ÉLÉMENTS DE PLANIFICATION")
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = 'Times New Roman'
    planning_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    planning_title.paragraph_format.space_after = Pt(4)

    planning_fields = [
        ("COMPÉTENCES DISCIPLINAIRES", parsed["competences_disc"]),
        ("COMP. TRANSV/COMP. TRANSDISC.", parsed["comp_transv"]),
        ("CONNAISSANCES TECHNIQUES", parsed["connaissances_tech"]),
        ("STRATÉGIE D'ENS/APPRENT./EVAL", parsed["strategie"]),
        ("MATÉRIEL", parsed["materiel"]),
    ]
    for label, value in planning_fields:
        p = doc.add_paragraph()
        run = p.add_run(f"{label} : ")
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        run = p.add_run(value or "[À compléter]")
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)

    # DEROULEMENT section
    doc.add_paragraph()
    derou_title = doc.add_paragraph()
    derou_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = derou_title.add_run("DÉROULEMENT")
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'

    # Create table
    rows = parsed["deroulement"]
    if rows:
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        # Header row
        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(["CONSIGNES", "RÉSULTATS ATTENDUS"]):
            hdr_cells[idx].text = ""
            p = hdr_cells[idx].paragraphs[0]
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_shading(hdr_cells[idx], "D9E2F3")

        # Set column widths
        for cell in table.columns[0].cells:
            cell.width = Cm(11)
        for cell in table.columns[1].cells:
            cell.width = Cm(6)

        # Data rows
        for row_data in rows:
            row = table.add_row()
            is_header = row_data.get("is_header", False)

            consigne_cell = row.cells[0]
            resultat_cell = row.cells[1]

            consigne_cell.text = ""
            resultat_cell.text = ""

            p = consigne_cell.paragraphs[0]
            run = p.add_run(row_data["consigne"])
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            if is_header:
                run.bold = True
                run.font.size = Pt(10)
                set_cell_shading(consigne_cell, "E8F0FE")
                set_cell_shading(resultat_cell, "E8F0FE")

            p = resultat_cell.paragraphs[0]
            run = p.add_run(row_data.get("resultat", ""))
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            if is_header:
                run.bold = True

    # Résumé
    if parsed["resume"]:
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run("RÉSUMÉ :")
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'

        p = doc.add_paragraph()
        run = p.add_run(parsed["resume"])
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'

    # Notes
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("NOTES PERSONNELLES")
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'
    if parsed["notes"]:
        p = doc.add_paragraph()
        run = p.add_run(parsed["notes"])
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'


def _generate_raw_docx(doc, content):
    """Fallback: generate document from raw markdown-like content."""
    lines = content.strip().split("\n")
    for line in lines:
        line_stripped = line.strip()

        if not line_stripped:
            doc.add_paragraph()
            continue

        # Headers
        if line_stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line_stripped[2:])
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = 'Times New Roman'
        elif line_stripped.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(line_stripped[3:])
            run.bold = True
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'
        elif line_stripped.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(line_stripped[4:])
            run.bold = True
            run.font.size = Pt(11)
            run.font.name = 'Times New Roman'
        elif line_stripped.startswith("**") and line_stripped.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line_stripped.strip("*"))
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'
        elif "|" in line_stripped:
            # Skip table separators
            if set(line_stripped.replace(" ", "")) <= set("|:-"):
                continue
            parts = [p.strip() for p in line_stripped.split("|")]
            parts = [p for p in parts if p]
            p = doc.add_paragraph()
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                run.font.size = Pt(10)
                run.font.name = 'Times New Roman'
                if idx < len(parts) - 1:
                    run = p.add_run("  |  ")
                    run.font.size = Pt(10)
        else:
            p = doc.add_paragraph()
            # Handle inline bold
            bold_parts = re.split(r'\*\*(.*?)\*\*', line_stripped)
            for idx, part in enumerate(bold_parts):
                run = p.add_run(part)
                run.font.size = Pt(10)
                run.font.name = 'Times New Roman'
                if idx % 2 == 1:
                    run.bold = True


def generate_pdf_from_docx(docx_path: str, pdf_path: str) -> str:
    """Convert DOCX to PDF using LibreOffice."""
    output_dir = os.path.dirname(pdf_path)
    try:
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", output_dir, docx_path],
            capture_output=True, timeout=60
        )
        # LibreOffice names the output based on input filename
        expected = os.path.join(output_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
        if os.path.exists(expected) and expected != pdf_path:
            os.rename(expected, pdf_path)
        return pdf_path
    except Exception as e:
        raise RuntimeError(f"Erreur conversion PDF: {e}")
